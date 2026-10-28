"""
Document Parser Service for Skill2Role RAG System.
Handles robust extraction and structural section parsing from PDF and DOCX files.
"""

import io
import re
import hashlib
from typing import Dict, Any, List, Optional
import pymupdf
import docx

KNOWN_SECTIONS = {
    "summary": ["summary", "professional summary", "profile", "about me", "objective", "career objective"],
    "experience": ["experience", "work experience", "employment history", "professional experience", "work history"],
    "skills": ["skills", "technical skills", "core competencies", "competencies", "technologies", "tech stack"],
    "projects": ["projects", "personal projects", "academic projects", "key projects"],
    "education": ["education", "academic background", "qualifications", "academic qualifications"],
    "certifications": ["certifications", "licenses", "courses", "certificates"]
}

def compute_content_hash(file_bytes: bytes) -> str:
    """Computes SHA-256 hash for document content deduplication."""
    return hashlib.sha256(file_bytes).hexdigest()

def classify_section_header(line: str) -> Optional[str]:
    """Matches a line against known resume section headers."""
    clean = re.sub(r"[^a-zA-Z\s]", "", line).strip().lower()
    if len(clean) < 3 or len(clean) > 35:
        return None
    for sec_type, aliases in KNOWN_SECTIONS.items():
        if clean in aliases:
            return sec_type
    return None

def parse_pdf(file_bytes: bytes) -> Dict[str, Any]:
    """
    Parses a PDF resume using PyMuPDF (fitz).
    Preserves text blocks, lines, and attempts section header detection.
    """
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    full_text = []
    lines_with_meta = []
    
    current_line_no = 1
    for page_idx in range(len(doc)):
        page = doc[page_idx]
        blocks = page.get_text("blocks")
        for b in blocks:
            # b: (x0, y0, x1, y1, text, block_no, block_type)
            block_text = b[4].strip()
            if not block_text:
                continue
            for line in block_text.splitlines():
                stripped = line.strip()
                if stripped:
                    lines_with_meta.append({
                        "line_no": current_line_no,
                        "text": stripped,
                        "page": page_idx + 1
                    })
                    current_line_no += 1
            full_text.append(block_text)

    raw_text = "\n".join(full_text)
    sections = segment_lines_into_sections(lines_with_meta)
    
    return {
        "raw_text": raw_text,
        "line_count": len(lines_with_meta),
        "lines": lines_with_meta,
        "sections": sections
    }

def parse_docx(file_bytes: bytes) -> Dict[str, Any]:
    """
    Parses a Word (.docx) resume using python-docx.
    Extracts paragraphs, runs, and table cells.
    """
    doc = docx.Document(io.BytesIO(file_bytes))
    lines_with_meta = []
    full_text = []
    current_line_no = 1

    for p in doc.paragraphs:
        stripped = p.text.strip()
        if stripped:
            lines_with_meta.append({
                "line_no": current_line_no,
                "text": stripped,
                "style": p.style.name if p.style else ""
            })
            current_line_no += 1
            full_text.append(stripped)

    # Also extract tables if present
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                line_str = " | ".join(row_texts)
                lines_with_meta.append({
                    "line_no": current_line_no,
                    "text": line_str,
                    "style": "Table"
                })
                current_line_no += 1
                full_text.append(line_str)

    raw_text = "\n".join(full_text)
    sections = segment_lines_into_sections(lines_with_meta)

    return {
        "raw_text": raw_text,
        "line_count": len(lines_with_meta),
        "lines": lines_with_meta,
        "sections": sections
    }

def segment_lines_into_sections(lines: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Groups ordered lines into named semantic sections.
    """
    sections = []
    current_sec_type = "header"
    current_sec_title = "Contact Information"
    current_lines = []

    for item in lines:
        text = item["text"]
        detected = classify_section_header(text)
        if detected:
            if current_lines:
                sections.append({
                    "section_type": current_sec_type,
                    "section_title": current_sec_title,
                    "text": "\n".join([l["text"] for l in current_lines]),
                    "lines": current_lines
                })
            current_sec_type = detected
            current_sec_title = text
            current_lines = []
        else:
            current_lines.append(item)

    if current_lines:
        sections.append({
            "section_type": current_sec_type,
            "section_title": current_sec_title,
            "text": "\n".join([l["text"] for l in current_lines]),
            "lines": current_lines
        })

    return sections

def parse_document(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Dispatcher based on filename extension.
    """
    fn = filename.lower()
    content_hash = compute_content_hash(file_bytes)
    
    if fn.endswith(".pdf"):
        parsed = parse_pdf(file_bytes)
        parsed["format"] = "pdf"
    elif fn.endswith(".docx"):
        parsed = parse_docx(file_bytes)
        parsed["format"] = "docx"
    else:
        # Fallback raw text parsing
        raw_text = file_bytes.decode("utf-8", errors="ignore")
        lines = [{"line_no": i+1, "text": l.strip()} for i, l in enumerate(raw_text.splitlines()) if l.strip()]
        parsed = {
            "raw_text": raw_text,
            "line_count": len(lines),
            "lines": lines,
            "sections": segment_lines_into_sections(lines),
            "format": "txt"
        }

    parsed["content_hash"] = content_hash
    parsed["filename"] = filename
    return parsed
