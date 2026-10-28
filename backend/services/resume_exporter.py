"""
Resume Exporter Service for Skill2Role RAG System.
Renders tailored resumes into professional ATS-friendly .docx and .pdf documents.
"""

import io
import re
import docx
from typing import Dict, Any, List
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from xml.sax.saxutils import escape as xml_escape

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

PLACEHOLDER_EMAIL_PATTERN = re.compile(r"@example\.(com|org|net)|@test\.(com|org)|test_candidate|dummy|mock", re.IGNORECASE)
PLACEHOLDER_PHONE_PATTERN = re.compile(r"9876543210|1234567890|\+?1?\s*555-", re.IGNORECASE)
PLACEHOLDER_LINKEDIN_PATTERN = re.compile(r"kushal-lead-dev|example-user|test-user|dummy-profile", re.IGNORECASE)

def detect_placeholder_contact_info(resume_data: Dict[str, Any]) -> List[str]:
    """
    Detects whether resume data contains synthetic/placeholder contact information.
    Returns a list of detected placeholder descriptions.
    """
    placeholders = []
    contact_info = resume_data.get("contact_info") or {}
    email = str(contact_info.get("email") or resume_data.get("candidate_email", ""))
    phone = str(contact_info.get("phone") or resume_data.get("candidate_phone", ""))
    linkedin = str(contact_info.get("linkedin") or resume_data.get("candidate_linkedin", ""))
    name = str(resume_data.get("candidate_name", ""))

    if email and PLACEHOLDER_EMAIL_PATTERN.search(email):
        placeholders.append(f"Synthetic email domain: '{email}'")
    if phone and PLACEHOLDER_PHONE_PATTERN.search(phone):
        placeholders.append(f"Synthetic dummy phone number: '{phone}'")
    if linkedin and PLACEHOLDER_LINKEDIN_PATTERN.search(linkedin):
        placeholders.append(f"Synthetic dummy LinkedIn handle: '{linkedin}'")
    if "test" in name.lower() and "candidate" in name.lower():
        placeholders.append(f"Synthetic test candidate name: '{name}'")
    return placeholders

def generate_docx(resume_data: Dict[str, Any], production_mode: bool = False) -> io.BytesIO:
    """
    Renders structured tailored resume into a clean ATS-compliant DOCX file.
    Includes contact info, summary, competencies, experience, education, and optional cover letter.
    If production_mode is True, blocks export if synthetic/placeholder contact information is detected.
    """
    placeholders = detect_placeholder_contact_info(resume_data)
    if production_mode and placeholders:
        raise ValueError(
            "PRODUCTION EXPORT BLOCKED: Resume contains synthetic or placeholder contact data: "
            + "; ".join(placeholders)
            + ". Real master resume and verified contact details must be provided before production export."
        )

    doc = docx.Document()

    # Set 0.75 inch margins for standard ATS format
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    if placeholders:
        banner = doc.add_paragraph()
        banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
        brun = banner.add_run("[AUDIT DEMO DATA — UNVERIFIED CONTACT DETAILS — DO NOT SUBMIT EXTERNALLY]")
        brun.bold = True
        brun.font.size = Pt(8.5)
        brun.font.color.rgb = RGBColor(220, 38, 38)

    name = str(resume_data.get("candidate_name") or "Candidate")
    target_role = str(resume_data.get("target_role") or "")
    target_company = str(resume_data.get("target_company") or "")
    summary = str(resume_data.get("tailored_summary") or "")
    matched_skills = resume_data.get("matched_skills") or []
    experiences = resume_data.get("experiences") or []
    education = resume_data.get("education") or []

    # Contact info
    contact_info = resume_data.get("contact_info") or {}
    email = contact_info.get("email") or resume_data.get("candidate_email", "")
    phone = contact_info.get("phone") or resume_data.get("candidate_phone", "")
    loc = contact_info.get("location") or resume_data.get("candidate_location", "")
    linkedin = contact_info.get("linkedin") or resume_data.get("candidate_linkedin", "")
    contact_parts = [p for p in [email, phone, loc, linkedin] if p]

    # Header: Name
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(name)
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(30, 41, 59)

    # Subtitle: Target Role & Company
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(f"{target_role} | Tailored for {target_company}")
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    # Contact line
    if contact_parts:
        c_p = doc.add_paragraph()
        c_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_run = c_p.add_run(" • ".join(contact_parts))
        c_run.font.size = Pt(9.5)
        c_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Professional Summary
    h1 = doc.add_paragraph()
    h1_run = h1.add_run("PROFESSIONAL SUMMARY")
    h1_run.bold = True
    h1_run.font.size = Pt(12)
    h1_run.font.color.rgb = RGBColor(15, 23, 42)

    sum_p = doc.add_paragraph(summary)
    sum_p.paragraph_format.space_after = Pt(10)

    # Core Competencies / Matched Skills
    if matched_skills:
        h_skills = doc.add_paragraph()
        hs_run = h_skills.add_run("CORE COMPETENCIES")
        hs_run.bold = True
        hs_run.font.size = Pt(12)
        hs_run.font.color.rgb = RGBColor(15, 23, 42)

        skills_p = doc.add_paragraph(", ".join(matched_skills))
        skills_p.paragraph_format.space_after = Pt(10)

    # Experience
    if experiences:
        h_exp = doc.add_paragraph()
        he_run = h_exp.add_run("PROFESSIONAL EXPERIENCE")
        he_run.bold = True
        he_run.font.size = Pt(12)
        he_run.font.color.rgb = RGBColor(15, 23, 42)

        for exp in experiences:
            exp_header = doc.add_paragraph()
            r_run = exp_header.add_run(exp.get("role", "Engineer"))
            r_run.bold = True
            r_run.font.size = Pt(11)

            c_run = exp_header.add_run(f" — {exp.get('company', '')}")
            c_run.font.size = Pt(11)

            p_run = exp_header.add_run(f" ({exp.get('period', '')})")
            p_run.italic = True
            p_run.font.size = Pt(10)
            p_run.font.color.rgb = RGBColor(100, 116, 139)

            for b in exp.get("bullets", []):
                raw_b_text = b.get("bullet_text", "") if isinstance(b, dict) else str(b)
                clean_b_text = re.sub(r"^[\s•\-\*–·]+", "", raw_b_text).strip()
                bp = doc.add_paragraph(clean_b_text, style="List Bullet")
                bp.paragraph_format.space_after = Pt(3)

            doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Education Section
    if education:
        h_edu = doc.add_paragraph()
        he_edu = h_edu.add_run("EDUCATION")
        he_edu.bold = True
        he_edu.font.size = Pt(12)
        he_edu.font.color.rgb = RGBColor(15, 23, 42)

        for edu in education:
            ep = doc.add_paragraph()
            d_run = ep.add_run(edu.get("degree", "Degree"))
            d_run.bold = True
            d_run.font.size = Pt(10.5)
            ins = edu.get("institution", "")
            if ins:
                ep.add_run(f", {ins}")
            yr = edu.get("year", "")
            if yr:
                ep.add_run(f" ({yr})")
            ep.paragraph_format.space_after = Pt(4)

    # Cover Letter Section
    cover_letter = resume_data.get("cover_letter", "")
    if cover_letter:
        doc.add_page_break()
        cl_h = doc.add_paragraph()
        cl_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cl_run = cl_h.add_run("COVER LETTER")
        cl_run.bold = True
        cl_run.font.size = Pt(14)
        
        for p in cover_letter.split("\n\n"):
            if p.strip():
                doc.add_paragraph(p.strip())

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

def generate_pdf(resume_data: Dict[str, Any], production_mode: bool = False) -> io.BytesIO:
    """
    Renders structured tailored resume into a clean ATS-compliant PDF using ReportLab.
    Safely escapes XML entities to prevent expat parsing errors.
    If production_mode is True, blocks export if synthetic/placeholder contact information is detected.
    """
    placeholders = detect_placeholder_contact_info(resume_data)
    if production_mode and placeholders:
        raise ValueError(
            "PRODUCTION EXPORT BLOCKED: Resume contains synthetic or placeholder contact data: "
            + "; ".join(placeholders)
            + ". Real master resume and verified contact details must be provided before production export."
        )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    # Custom Clean Styles
    title_style = ParagraphStyle(
        "ResumeTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        alignment=1, # Center
        textColor=colors.HexColor("#1e293b")
    )

    sub_style = ParagraphStyle(
        "ResumeSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        alignment=1,
        textColor=colors.HexColor("#64748b")
    )

    contact_style = ParagraphStyle(
        "ResumeContact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        alignment=1,
        textColor=colors.HexColor("#64748b")
    )

    section_heading = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#0f172a"),
        spaceBefore=10,
        spaceAfter=4
    )

    body_style = ParagraphStyle(
        "ResumeBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#334155")
    )

    bullet_style = ParagraphStyle(
        "ResumeBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        leftIndent=15,
        textColor=colors.HexColor("#334155"),
        spaceAfter=3
    )

    story = []

    if placeholders:
        banner_style = ParagraphStyle(
            "DemoNotice",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            alignment=1,
            textColor=colors.HexColor("#dc2626"),
            spaceAfter=6
        )
        story.append(Paragraph("[AUDIT DEMO DATA — UNVERIFIED CONTACT DETAILS — DO NOT SUBMIT EXTERNALLY]", banner_style))

    name = str(resume_data.get("candidate_name") or "Candidate")
    target_role = str(resume_data.get("target_role") or "")
    target_company = str(resume_data.get("target_company") or "")
    summary = str(resume_data.get("tailored_summary") or "")
    matched_skills = resume_data.get("matched_skills") or []
    experiences = resume_data.get("experiences") or []
    education = resume_data.get("education") or []

    # Contact info
    contact_info = resume_data.get("contact_info") or {}
    email = contact_info.get("email") or resume_data.get("candidate_email", "")
    phone = contact_info.get("phone") or resume_data.get("candidate_phone", "")
    loc = contact_info.get("location") or resume_data.get("candidate_location", "")
    linkedin = contact_info.get("linkedin") or resume_data.get("candidate_linkedin", "")
    contact_parts = [xml_escape(p) for p in [email, phone, loc, linkedin] if p]

    story.append(Paragraph(xml_escape(name), title_style))
    story.append(Paragraph(f"{xml_escape(target_role)} | Tailored for {xml_escape(target_company)}", sub_style))
    if contact_parts:
        story.append(Paragraph(" • ".join(contact_parts), contact_style))
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=10))

    # Summary
    story.append(Paragraph("PROFESSIONAL SUMMARY", section_heading))
    story.append(Paragraph(xml_escape(summary), body_style))
    story.append(Spacer(1, 8))

    # Skills
    if matched_skills:
        story.append(Paragraph("CORE COMPETENCIES", section_heading))
        escaped_skills = [xml_escape(s) for s in matched_skills]
        story.append(Paragraph(", ".join(escaped_skills), body_style))
        story.append(Spacer(1, 8))

    # Experience
    if experiences:
        story.append(Paragraph("PROFESSIONAL EXPERIENCE", section_heading))
        for exp in experiences:
            role_esc = xml_escape(exp.get("role", ""))
            comp_esc = xml_escape(exp.get("company", ""))
            period_esc = xml_escape(exp.get("period", ""))
            role_header = f"<b>{role_esc}</b> — {comp_esc} <font color='#64748b'>({period_esc})</font>"
            story.append(Paragraph(role_header, body_style))
            for b in exp.get("bullets", []):
                raw_b_text = b.get("bullet_text", "") if isinstance(b, dict) else str(b)
                clean_b_text = re.sub(r"^[\s•\-\*–·]+", "", raw_b_text).strip()
                story.append(Paragraph(f"• {xml_escape(clean_b_text)}", bullet_style))
            story.append(Spacer(1, 6))

    # Education Section
    if education:
        story.append(Paragraph("EDUCATION", section_heading))
        for edu in education:
            deg_esc = xml_escape(edu.get("degree", ""))
            ins_esc = xml_escape(edu.get("institution", ""))
            yr_esc = xml_escape(str(edu.get("year", "")))
            edu_line = f"<b>{deg_esc}</b>"
            if ins_esc:
                edu_line += f", {ins_esc}"
            if yr_esc:
                edu_line += f" <font color='#64748b'>({yr_esc})</font>"
            story.append(Paragraph(edu_line, body_style))
            story.append(Spacer(1, 4))

    doc.build(story)
    buf.seek(0)
    return buf
